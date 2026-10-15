"""Blind second-annotator re-check of the deciding-clause label for ALL 94 in-scope, named-institution questions.
Reason: an automated check (analysis/retrieval_misses.py) found 12 questions labelled NO whose cited clause appears in the frozen
retrieved text; the round-1 recheck verified quotes for YES only. To avoid one-directional correction, every question is re-checked,
blind to the original label, the automated flags and all answers. Same rule as the round-1 recheck; quote required for YES/PARTIAL.
Outputs labelling/clause_recheck2/clause_recheck2.xlsx and CLAUSE_RECHECK2_KEY_do_not_share.csv."""
import csv, json, os, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__)); RAG = os.path.dirname(HERE)
OUT = os.path.join(HERE, "clause_recheck2"); os.makedirs(OUT, exist_ok=True)
key = {r["qid"]: r for r in csv.DictReader(open(os.path.join(RAG, "answer_key_FROZEN_2026-09-30.csv"), encoding="utf-8"))}
lab = {r["qid"]: r for r in csv.DictReader(open(os.path.join(RAG, "labels_FROZEN_2026-09-30.csv"), encoding="utf-8"))}
fr = {json.loads(l)["qid"]: json.loads(l)["chunk_ids"] for l in open(os.path.join(RAG, "index", "retrieval_frozen.jsonl"), encoding="utf-8")}
ch = {json.loads(l)["chunk_id"]: json.loads(l) for l in open(os.path.join(RAG, "index", "chunks.jsonl"), encoding="utf-8")}
qs = sorted(q for q, r in lab.items() if r["clause_in_retrieved"] in ("YES", "PARTIAL", "NO"))
assert len(qs) == 94
random.Random(20261002).shuffle(qs)
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); g = wb.active; g.title = "How to use"
HOW = ["Clause re-check (94 questions)", "",
       "For each row, answer one question: does the Retrieved policy text contain the clause needed to answer the employee's question correctly?",
       "Judge ONLY the text shown in 'Retrieved policy text'. Use the 'Reference clauses' column to know which rule decides the question.",
       "YES = the deciding sentence is present, including the number or condition it turns on.",
       "PARTIAL = part of the deciding rule is present, but its decisive number, condition or qualifier is missing.",
       "NO = the deciding rule is absent. Source lines, headers, definitions, or a different rule do not count.",
       "For every YES or PARTIAL, copy-paste the deciding sentence (or the part present) from 'Retrieved policy text' into the Quote column.",
       "For every NO, write in Notes which part of the reference clause you looked for and did not find (one short phrase).",
       "Work alone, no AI tools. Only edit the yellow columns. About 20-30 seconds per row."]
for i, t in enumerate(HOW, 1):
    c = g.cell(row=i, column=1, value=t); c.font = F(bold=(i == 1), size=14 if i == 1 else 10); c.alignment = Alignment(wrap_text=True)
g.column_dimensions["A"].width = 130
ws = wb.create_sheet("Clause")
cols = [("row_id", 8), ("Institution", 16), ("Question", 45), ("Retrieved policy text", 75), ("Gold verdict", 14), ("Reference clauses", 45),
        ("Clause present", 14), ("Quote from retrieved text", 45), ("Notes", 35)]
for j, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    c.border = B; c.alignment = Alignment(wrap_text=True); ws.column_dimensions[c.column_letter].width = w
krows = []
for i, q in enumerate(qs, 2):
    k = key[q]; ctx = "\n\n".join(f"[{ch[c]['doc']}] {ch[c]['text']}" for c in fr[q])
    vals = [f"K{i - 1:03d}", k["institution"], k["question"], ctx, k["gold_verdict"], f"A: {k['clause_annotator_A']}\nB: {k['clause_annotator_B']}", "", "", ""]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = F(size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical="top")
        if j >= 7: c.fill = PatternFill("solid", fgColor="FFFF00")
    krows.append({"row_id": f"K{i - 1:03d}", "qid": q})
dv = DataValidation(type="list", formula1='"YES,PARTIAL,NO"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"G2:G{len(qs) + 1}")
ws.freeze_panes = "D2"; wb.active = 1
wb.save(os.path.join(OUT, "clause_recheck2.xlsx"))
with open(os.path.join(OUT, "CLAUSE_RECHECK2_KEY_do_not_share.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=["row_id", "qid"]); w.writeheader(); w.writerows(krows)
print("built", len(qs), "rows")
