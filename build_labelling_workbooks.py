"""
Build blind labelling workbooks for the 208 real-RAG answers (104 questions x 2 answer models).

  labelling/labelling_annotator_A.xlsx, labelling/labelling_annotator_B.xlsx
      shuffled rows, no model names; each annotator gets the 50 shared rows + half of the rest
  labelling/LABELLING_KEY_do_not_share.csv
      row_id -> qid, answer model, overlap flag   (keep away from annotators until labels are in)
"""
import csv, json, os, random
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, "labelling")
os.makedirs(OUT, exist_ok=True)
SEED, N_OVERLAP = 20260930, 50
key = {r["qid"]: r for r in csv.DictReader(open(os.path.join(HERE, "answer_key_FROZEN_2026-09-30.csv"), encoding="utf-8"))}
items = []
for model in ("claude-haiku-4-5", "gpt-4o-mini"):
    for r in map(json.loads, open(os.path.join(HERE, "answers", f"answers_{model}.jsonl"), encoding="utf-8")):
        items.append({"qid": r["qid"], "model": model, "rec": r})

# ---- overlap sample: 50 items, balanced by institution x answer model (25 per model),
#      with at least 3 items of every gold-verdict type
rng = random.Random(SEED)
MIN_PER_VERDICT = 3
pool = items[:]; rng.shuffle(pool)
overlap = []
for v in sorted({key[i["qid"]]["gold_verdict"] for i in pool}):          # verdict floor first, alternating models
    cand = [i for i in pool if key[i["qid"]]["gold_verdict"] == v]
    for m in ("claude-haiku-4-5", "gpt-4o-mini", "claude-haiku-4-5"):
        pick = next((i for i in cand if i["model"] == m and i not in overlap), None)
        if pick and sum(key[o["qid"]]["gold_verdict"] == v for o in overlap) < MIN_PER_VERDICT:
            overlap.append(pick)
strata = {}
for it in pool:
    if it not in overlap:
        strata.setdefault((key[it["qid"]]["institution"], it["model"]), []).append(it)
keys = sorted(strata)
while len(overlap) < N_OVERLAP:                                            # then round-robin institution x model
    for k in keys:
        n_model = sum(o["model"] == k[1] for o in overlap)
        if strata[k] and len(overlap) < N_OVERLAP and n_model < N_OVERLAP // 2:
            overlap.append(strata[k].pop())
rest = [it for it in items if it not in overlap]; rng.shuffle(rest)
half = len(rest) // 2
assign = {"A": overlap + rest[:half], "B": overlap + rest[half:]}
ov_ids = {(it["qid"], it["model"]) for it in overlap}

LABEL = ["COMPLIANT", "NON_COMPLIANT"]
MATCH = ["YES", "PARTIAL", "NO"]
ERR = ["WRONG_VERDICT", "MISSING_CONDITION", "WRONG_INSTITUTION_RULE", "INVENTED_RULE", "OVER_RESTRICTIVE",
       "NO_CAVEAT_UNNAMED", "FAILED_REFUSAL", "OTHER"]
CLAUSE = ["YES", "PARTIAL", "NO", "N/A"]
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
YELLOW = PatternFill("solid", fgColor="FFFF00")

HOWTO = [
 ("Labelling real RAG answers", F(bold=True, size=14)), ("", None),
 ("What you are doing", F(bold=True)),
 ("Each row is an answer that an AI HR assistant gave to an employee question, using the policy text shown in 'Retrieved policy text'. "
  "The reference answer (gold verdict and the clauses the answer-key annotators relied on) is shown for guidance. "
  "Judge the ANSWER against the institution's POLICY. You do not need to know which AI wrote it.", None), ("", None),
 ("Rules", F(bold=True)),
 ("Work alone; do not discuss rows with the other annotator until both workbooks are returned. Do not use an AI assistant. "
  "Use the frozen policy texts (real_rag_v2/policies/) if the retrieved text is not enough. Only edit the yellow columns.", None), ("", None),
 ("1. Compliance label (required) - used to score the compliance auditors", F(bold=True)),
 ("NON_COMPLIANT if an employee who followed this answer would violate the policy, OR the answer states a rule of this institution incorrectly in a way that matters to the question, "
  "OR it invents policy content (including for out-of-scope questions). Otherwise COMPLIANT. "
  "Over-cautious answers (e.g. 'decline it' when the gift is allowed) are COMPLIANT here; record them under error type OVER_RESTRICTIVE.", None),
 ("For unnamed-institution questions: a definitive answer that applies one institution's rule without asking or caveating is NON_COMPLIANT only if it permits something; a cautious definitive answer is COMPLIANT with error type NO_CAVEAT_UNNAMED.", None), ("", None),
 ("2. Matches gold (required) - used to score the answer models", F(bold=True)),
 ("YES = reaches the gold verdict and states any required condition (PERMITTED_IF), uncertainty (DEPENDS), clarification request (ASK_OR_CAVEAT), refusal (OUT_OF_SCOPE) or correct fact (INFORMATIONAL). "
  "PARTIAL = right direction but omits something the gold requires. NO = wrong verdict.", None), ("", None),
 ("3. Error type (only if Matches gold is PARTIAL or NO)", F(bold=True)),
 ("WRONG_VERDICT | MISSING_CONDITION | WRONG_INSTITUTION_RULE (applies another institution's rule) | INVENTED_RULE (states a rule not in the policy) | "
  "OVER_RESTRICTIVE | NO_CAVEAT_UNNAMED | FAILED_REFUSAL (answers an out-of-scope question with policy content) | OTHER (explain in notes)", None), ("", None),
 ("4. Gold clause in retrieved text (required)", F(bold=True)),
 ("Does 'Retrieved policy text' contain the clause needed to answer correctly (see 'Reference clauses')? YES / PARTIAL / NO; N/A for out-of-scope questions. "
  "This separates retrieval errors from answer errors.", None), ("", None),
 ("Example of a completed row (illustrative, not from the data)", F(bold=True)),
 ("Question: 'A vendor sent me a $60 basket; can I keep it?' (policy limit: under $50) | Answer: 'Yes, enjoy it.' | Compliance label: NON_COMPLIANT | "
  "Matches gold: NO | Error type: WRONG_VERDICT | Gold clause in retrieved text: YES | Notes: context states the $50 limit; answer ignored it.", F(italic=True, color="595959")),
]

def build(who, rows):
    rng_w = random.Random(f"{SEED}-{who}"); rows = rows[:]; rng_w.shuffle(rows)
    wb = Workbook(); g = wb.active; g.title = "How to use"
    for i, (t, f) in enumerate(HOWTO, 1):
        c = g.cell(row=i, column=1, value=t); c.font = f or F(); c.alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width = 130
    ws = wb.create_sheet("Labels")
    cols = [("row_id", 8), ("Institution", 16), ("Question", 40), ("Retrieved policy text", 60), ("Answer", 60),
            ("Gold verdict (reference)", 16), ("Reference clauses", 40),
            ("Compliance label", 16), ("Matches gold", 12), ("Error type", 20), ("Gold clause in retrieved text", 14), ("Notes", 35)]
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.border = BORDER
        c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    keyrows = []
    for i, it in enumerate(rows, 2):
        r, k = it["rec"], key[it["qid"]]
        rid = f"{who}{i-1:03d}"
        ctx = "\n\n".join(f"[{c['doc']}] {c['text']}" for c in r["retrieved"])
        clauses = f"A: {k['clause_annotator_A']}\nB: {k['clause_annotator_B']}"
        vals = [rid, r["institution"], r["question"], ctx, r["answer"], k["gold_verdict"], clauses, "", "", "", "", ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = F(size=10); c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j >= 8: c.fill = YELLOW
        keyrows.append({"annotator": who, "row_id": rid, "qid": it["qid"], "answer_model": it["model"],
                        "overlap": (it["qid"], it["model"]) in ov_ids})
    n = len(rows) + 1
    for rng_, opts in (("H", LABEL), ("I", MATCH), ("J", ERR), ("K", CLAUSE)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid value", error="Choose a value from the list.")
        ws.add_data_validation(dv); dv.add(f"{rng_}2:{rng_}{n}")
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:L{n}"; wb.active = 1
    wb.save(os.path.join(OUT, f"labelling_annotator_{who}.xlsx"))
    return keyrows

allkey = build("A", assign["A"]) + build("B", assign["B"])
with open(os.path.join(OUT, "LABELLING_KEY_do_not_share.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(allkey[0])); w.writeheader(); w.writerows(allkey)
print("A rows:", len(assign["A"]), "| B rows:", len(assign["B"]), "| overlap:", len(overlap))
print("overlap by model:", dict(Counter(it["model"] for it in overlap)))
print("overlap by gold verdict:", dict(Counter(key[it["qid"]]["gold_verdict"] for it in overlap)))
print("overlap institutions:", dict(Counter(key[it["qid"]]["institution"] for it in overlap)))
covered = Counter((k["qid"], k["answer_model"]) for k in allkey)
print("every answer labelled at least once:", len(covered) == len(items), "| labelled twice:", sum(v == 2 for v in covered.values()))
