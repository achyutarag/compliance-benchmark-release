"""Assemble and freeze the final human labels for all 208 real-RAG answers."""
import csv, datetime, hashlib, json, os
from collections import Counter
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__)); L = os.path.join(HERE, "labelling"); R = os.path.join(L, "reconciliation")
lk = {r["row_id"]: r for r in csv.DictReader(open(os.path.join(L, "LABELLING_KEY_do_not_share.csv"), encoding="utf-8"))}
gold = {r["qid"]: r for r in csv.DictReader(open(os.path.join(HERE, "answer_key_FROZEN_2026-09-30.csv"), encoding="utf-8"))}
MAP = {"FABRICATED_RULE": "INVENTED_RULE", "WRONG_CLAUSE": "OTHER"}

def load(w):
    wb = load_workbook(os.path.join(L, "completed", f"labelling_annotator_{w}_completed.xlsx"), read_only=True)
    ws = next(s for s in wb.worksheets if next(s.iter_rows(values_only=True, max_row=1))[0] == "row_id")
    rows = list(ws.iter_rows(values_only=True)); h = rows[0]; out = {}
    for r in rows[1:]:
        if r[0]:
            d = {k: ("" if v is None else str(v).strip()) for k, v in zip(h, r)}
            d["Error type"] = MAP.get(d["Error type"], d["Error type"])
            out[(lk[d["row_id"]]["qid"], lk[d["row_id"]]["answer_model"])] = d
    return out
A, B = load("A"), load("B")
rw = load_workbook(os.path.join(R, "label_reconciliation.xlsx"), read_only=True)
recon = {}
for r in list(rw["Disagreements"].iter_rows(values_only=True))[1:]:
    a = r[0].split(" / ")[0]; recon[(lk[a]["qid"], lk[a]["answer_model"])] = (r[12], r[13], r[14] or "", r[15])
flags = {}
for r in list(rw["Consistency flags"].iter_rows(values_only=True))[1:]:
    flags[r[1]] = (r[9], r[10], r[11] or "", r[12])
ERR_FIX = {("A120", "B006"): ("OVER_RESTRICTIVE", "answer says do not accept; more cautious than DEPENDS requires"),
           ("A083", "B090"): ("OVER_RESTRICTIVE", "answer resolves DEPENDS as 'do not keep'"),
           ("A125", "B056"): ("INVENTED_RULE", "states that $45 exceeds the threshold (false)"),
           ("A114", "B103"): ("WRONG_VERDICT", "permits without institution caveat"),
           ("A089", "B045"): ("INVENTED_RULE", "claims the policy bans all gifts from external parties")}
clause = {(r["qid"], r["answer_model"]): r for r in csv.DictReader(open(os.path.join(R, "clause_recheck_validated.csv"), encoding="utf-8"))}
D10_QUOTE = ("Employees having authority or influence over matters regarding purchasing or contracting should exercise caution and good "
             "judgment in accepting gifts from individuals, entities, or key stakeholders of entities that are doing business or seeking to do "
             "business with the university, even if the gifts are of nominal value ($150 or less).")

final = []
for k in sorted(set(A) | set(B)):
    q, m = k; a, b = A.get(k), B.get(k)
    if a and b:
        if k in recon:
            comp, match, err, note = recon[k]; src = "shared: reconciled"
        else:
            comp, match, err = a["Compliance label"], a["Matches gold"], a["Error type"]; src = "shared: A and B agreed"
            note = "A: " + a["Notes"] + " | B: " + b["Notes"]
            fx = ERR_FIX.get((a["row_id"], b["row_id"]))
            if fx: err = fx[0]; src += "; error type resolved by written definitions"; note += f" | error type: {fx[1]}"
        rows_ids = f"{a['row_id']}/{b['row_id']}"
    else:
        d = a or b; who = "A" if a else "B"
        comp, match, err, note = d["Compliance label"], d["Matches gold"], d["Error type"], d["Notes"]; src = f"single annotator {who}"
        rows_ids = d["row_id"]
        if d["row_id"] in flags:
            comp, match, err, fnote = flags[d["row_id"]]; src += "; consistency-flag correction (written rule)"; note += " | " + fnote
    c = clause[k]; cl, quote, csrc = c["clause_label"], c["quote"], "clause re-check (single annotator, quote-verified)"
    if q == "D10" and cl == "NO":
        cl, quote, csrc = "YES", D10_QUOTE, "clause re-check; documented correction: deciding sentence present verbatim in chunk 2 (matches agreed example)"
    final.append({"qid": q, "answer_model": m, "institution": gold[q]["institution"], "structures": gold[q]["structures"],
                  "pair_group": gold[q]["pair_group"], "gold_verdict": gold[q]["gold_verdict"],
                  "compliance_label": comp, "y_violation": 1 if comp == "NON_COMPLIANT" else 0,
                  "matches_gold": match, "error_type": err if match in ("PARTIAL", "NO") else "",
                  "clause_in_retrieved": cl, "clause_quote": quote, "label_source": src, "clause_source": csrc,
                  "labelling_rows": rows_ids, "notes": note})
assert len(final) == 208 and all(f["compliance_label"] in ("COMPLIANT", "NON_COMPLIANT") for f in final)
assert all(f["matches_gold"] in ("YES", "PARTIAL", "NO") for f in final)
out = os.path.join(HERE, "labels_FROZEN_2026-09-30.csv")
with open(out, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(final[0])); w.writeheader(); w.writerows(final)
sha = hashlib.sha256(open(out, "rb").read()).hexdigest()
man = {"file": os.path.basename(out), "sha256": sha, "frozen_at": datetime.datetime.now().isoformat(timespec="seconds"),
       "n_answers": 208, "answer_models": dict(Counter(f["answer_model"] for f in final)),
       "compliance": dict(Counter(f["compliance_label"] for f in final)),
       "matches_gold": dict(Counter(f["matches_gold"] for f in final)),
       "error_type": dict(Counter(f["error_type"] or "(none)" for f in final)),
       "clause_in_retrieved": dict(Counter(f["clause_in_retrieved"] for f in final)),
       "label_source": dict(Counter(f["label_source"].split(";")[0] for f in final)),
       "inter_annotator_agreement_shared_rows_n50": {"compliance": {"raw": 0.90, "kappa": 0.56, "PABAK": 0.80, "positive_agreement": 0.62},
                                                     "matches_gold": {"raw": 0.80, "kappa": 0.59},
                                                     "clause_original_labels": {"raw": 0.68, "kappa": 0.46},
                                                     "error_type": {"raw": 0.70, "kappa": 0.44}},
       "reconciliation": {"shared_disagreements_resolved": 11, "consistency_flag_corrections": 3,
                          "error_type_resolutions_by_definition": 5, "error_type_mappings": MAP},
       "clause_recheck": {"annotator": "single", "quotes_verified_verbatim": "102/102", "agreed_examples_reproduced": "15/16",
                          "documented_corrections": ["D10 x2: NO -> YES (deciding sentence verbatim in retrieved chunk 2)"]},
       "depends_on": {"answer_key_sha256": "689571cb91e110968b02db51c5223434cbea2719790934c10d7fea928c494420",
                      "corpus_sha256": json.load(open(os.path.join(HERE, "policies", "MANIFEST.json"), encoding="utf-8"))["corpus_sha256"]}}
json.dump(man, open(os.path.join(HERE, "labels_FROZEN_MANIFEST.json"), "w", encoding="utf-8"), indent=2)
os.chmod(out, 0o444)
print("FROZEN", out, "\nsha256", sha)
for k in ("compliance", "matches_gold", "clause_in_retrieved", "label_source", "answer_models"): print(f"  {k}: {man[k]}")
