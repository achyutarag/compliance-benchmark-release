"""Freeze version 2 of the deciding-clause labels (94 in-scope named-institution questions) and recompute the clause-dependent tables.
Sources: 71 questions where round 1 and annotator 3 agreed; 9 auto-adopted (round-1 NO, annotator-3 YES with the cited clause verbatim in the
retrieved text); 14 adjudicated (returned sheet, 2026-10-11). Unnamed-institution and out-of-scope questions keep N/A."""
import csv, datetime, hashlib, json, os
from collections import Counter
import numpy as np
from scipy.stats import binomtest
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); CD = os.path.join(HERE, "labelling", "clause_recheck2")
OUT = os.path.join(HERE, "clause_labels_v2_FROZEN_2026-10-11.csv"); assert not os.path.exists(OUT), "already frozen"
key = {r["row_id"]: r["qid"] for r in csv.DictReader(open(os.path.join(CD, "CLAUSE_RECHECK2_KEY_do_not_share.csv"), encoding="utf-8"))}
a3 = {key[r["Row ID"]]: r for r in csv.DictReader(open(os.path.join(CD, "clause_recheck2_annotator3_returned.csv"), encoding="utf-8"))}
r1 = {}
for r in csv.DictReader(open(os.path.join(HERE, "labels_FROZEN_2026-09-30.csv"), encoding="utf-8")): r1.setdefault(r["qid"], r)
adj = {}
for row in list(load_workbook(os.path.join(CD, "adjudication_sheet_returned.xlsx"), read_only=True)["Adjudicate"].iter_rows(values_only=True))[1:]:
    adj[row[0]] = (row[10], row[11])
AUTO = {"C06", "H08", "T01", "T06", "U04", "U10", "U11", "U14", "V09"}
rows, src = [], Counter()
for q in sorted(r1):
    old = r1[q]["clause_in_retrieved"]
    if q in a3:
        if q in adj: lab, s, note = adj[q][0], "adjudicated", adj[q][1]
        elif q in AUTO: lab, s, note = a3[q]["Clause present"], "auto-adopted (verbatim check)", ""
        else:
            assert old == a3[q]["Clause present"]; lab, s, note = old, "round 1 and annotator 3 agree", ""
        quote = a3[q]["Quote from retrieved text"] if lab == a3[q]["Clause present"] else r1[q]["clause_quote"]
    else: lab, s, note, quote = old, "n/a (unnamed institution or out of scope)", "", ""
    src[s] += 1; rows.append({"qid": q, "clause_in_retrieved_v2": lab, "round1_label": old, "annotator3_label": a3[q]["Clause present"] if q in a3 else "",
                              "source": s, "quote": quote, "adjudication_note": note})
with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
os.chmod(OUT, 0o444)
V2 = {r["qid"]: r["clause_in_retrieved_v2"] for r in rows}
man = {"file": os.path.basename(OUT), "sha256": hashlib.sha256(open(OUT, "rb").read()).hexdigest(), "frozen_at": datetime.datetime.now().isoformat(timespec="seconds"),
       "sources": dict(src), "label_counts": dict(Counter(V2[q] for q in a3)),
       "agreement_before_adjudication": {"three_category_raw": "71/94", "three_category_kappa": 0.57, "yes_vs_not_yes_raw": "79/94", "yes_vs_not_yes_kappa": 0.69},
       "note": "Supersedes the clause_in_retrieved column of labels_FROZEN_2026-09-30.csv and labels_ext_FROZEN_2026-10-01.csv (compliance labels unchanged)."}
json.dump(man, open(os.path.join(HERE, "clause_labels_v2_FROZEN_MANIFEST.json"), "w"), indent=2)
print(json.dumps(man, indent=2))
print("changed vs round 1:", sum(1 for q in a3 if V2[q] != r1[q]["clause_in_retrieved"]), "of 94 | final equals annotator 3 on", sum(1 for q in a3 if V2[q] == a3[q]["Clause present"]), "of 94")

# ---- recompute clause-dependent numbers
A = {(r["qid"], r["answer_model"]): r for r in csv.DictReader(open(os.path.join(HERE, "results", "all416_per_answer.csv"), encoding="utf-8"))}
ORIG, EXT = ("claude-haiku-4-5", "gpt-4o-mini"), ("Llama-3.2-3B-Instruct", "gemma-2-2b-it")
print("\nVIOLATIONS BY CLAUSE PRESENCE (v2): questions | answers | violations Haiku+4o-mini | violations Llama+Gemma")
for c in ("YES", "PARTIAL", "NO", "N/A"):
    ks = [k for k in A if V2[k[0]] == c]
    print(f"  {c:8} {len({k[0] for k in ks}):2d} | {len(ks):3d} | {sum(int(A[k]['y']) for k in ks if k[1] in ORIG):2d} | {sum(int(A[k]['y']) for k in ks if k[1] in EXT):2d}")
print("\nJUDGE / AUDITOR FALSE-POSITIVE RATE (compliant answers) by clause presence, retrieved context:")
for col, nm in (("gpt55_retrieved", "GPT-5.5"), ("sonnet55_retrieved", "Sonnet 5.5"), ("mscg", "MSCG")):
    for c in ("YES", "NO"):
        ks = [k for k in A if V2[k[0]] == c and A[k]["y"] == "0"]
        print(f"  {nm:10} clause={c:3}: {sum(int(A[k][col]) for k in ks)}/{len(ks)} = {sum(int(A[k][col]) for k in ks) / len(ks):.0%}")
print("\nRECALL by clause presence (retrieved context):")
for c in ("YES", "PARTIAL", "NO", "N/A"):
    ks = [k for k in A if V2[k[0]] == c and A[k]["y"] == "1"]
    print(f"  clause={c:7} violations={len(ks):2d}: " + ", ".join(f"{nm} {sum(int(A[k][col]) for k in ks)}" for col, nm in (("mscg", "MSCG"), ("gpt55_retrieved", "GPT"), ("sonnet55_retrieved", "Sonnet"))))
ins = [q for q in a3]; print(f"\nin-scope questions with clause fully present: {sum(1 for q in ins if V2[q] == 'YES')} of {len(ins)}; partial {sum(1 for q in ins if V2[q] == 'PARTIAL')}; absent {sum(1 for q in ins if V2[q] == 'NO')}")
# stronger vs small models: share of violations with clause present
for g, ms in (("Haiku+4o-mini", ORIG), ("Llama+Gemma", EXT)):
    v = [k for k in A if k[1] in ms and A[k]["y"] == "1" and V2[k[0]] != "N/A"]
    print(f"  {g}: of {len(v)} in-scope violations, clause YES {sum(V2[k[0]] == 'YES' for k in v)}, PARTIAL {sum(V2[k[0]] == 'PARTIAL' for k in v)}, NO {sum(V2[k[0]] == 'NO' for k in v)}")
