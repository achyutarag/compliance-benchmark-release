"""Agreement between the LLM failure codes and annotator A's blind human codes (38 rows)."""
import csv, os, sys
from collections import Counter
from openpyxl import load_workbook
H = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1]
ws = load_workbook(src, data_only=True)["Coding"]
hdr = [c.value for c in ws[1]]; ci = hdr.index("How the answer went wrong"); ni = hdr.index("Notes")
human = {r[0].value: (r[ci].value or "").strip() for r in ws.iter_rows(min_row=2) if r[0].value}
key = {r["row_id"]: r for r in csv.DictReader(open(os.path.join(H, "labelling_ext", "validation", "FAILURE_CODE_KEY_do_not_share.csv"), encoding="utf-8"))}
assert set(human) == set(key), "row ids differ"
rows = [(k, key[k]["llm_code"], human[k]) for k in sorted(key)]
print("blank:", [k for k, _, h in rows if not h])
n = len(rows); agree = sum(l == h for _, l, h in rows); print(f"raw agreement {agree}/{n} = {agree/n:.0%}")
cats = sorted({l for _, l, _ in rows} | {h for _, _, h in rows})
pl = Counter(l for _, l, _ in rows); ph = Counter(h for _, _, h in rows)
pe = sum(pl[c] / n * ph[c] / n for c in cats); kappa = (agree / n - pe) / (1 - pe) if pe < 1 else float("nan")
print(f"Cohen kappa {kappa:.2f}"); print("LLM  :", dict(pl)); print("Human:", dict(ph))
print("confusion (LLM -> human):", dict(Counter((l, h) for _, l, h in rows)))
grp = lambda c: "prohibition_or_condition" if c in ("OVERRIDDEN_PROHIBITION", "MISAPPLIED_CONDITION") else c
ag2 = sum(grp(l) == grp(h) for _, l, h in rows); print(f"merged prohibition+condition agreement {ag2}/{n}")
print("disagreements:")
for r in ws.iter_rows(min_row=2):
    k = r[0].value
    if k and key[k]["llm_code"] != human[k]: print(" ", k, key[k]["llm_code"], "->", human[k], "|", (r[ni].value or "")[:150])
