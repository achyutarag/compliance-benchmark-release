"""
Round-2 (extension) reconciliation workbooks: every answer where A and B differ on compliance label or matches-gold.
  labelling_ext/reconciliation/reconciliation_ext_annotator_A.xlsx  (same rows for B; each annotator works alone)
Rows are sorted by question so both answers to one question sit together. Model names are never shown.
Each annotator re-reads the answer and either keeps or revises their own label, with a one-line reason.
"""
import csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # repo root; L = os.path.join(HERE, "labelling_ext")
OUT = os.path.join(L, "reconciliation"); os.makedirs(OUT, exist_ok=True)
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
YELLOW = PatternFill("solid", fgColor="FFFF00"); HEAD = PatternFill("solid", fgColor="305496"); GREY = PatternFill("solid", fgColor="EDEDED")
LABEL, MATCH = ["COMPLIANT", "NON_COMPLIANT"], ["YES", "PARTIAL", "NO"]
ERR = ["WRONG_VERDICT", "MISSING_CONDITION", "WRONG_INSTITUTION_RULE", "INVENTED_RULE", "OVER_RESTRICTIVE",
       "NO_CAVEAT_UNNAMED", "FAILED_REFUSAL", "OTHER"]

gold = {r["qid"]: r for r in csv.DictReader(open(os.path.join(HERE, "answer_key_FROZEN_2026-09-30.csv"), encoding="utf-8"))}
ans = {}
for m in ("Llama-3.2-3B-Instruct", "gemma-2-2b-it"):
    for r in map(json.loads, open(os.path.join(HERE, "answers", f"answers_{m}.jsonl"), encoding="utf-8")):
        ans[(r["qid"], m)] = r
D = list(csv.DictReader(open(os.path.join(L, "completed", "disagreements_ext.csv"), encoding="utf-8")))
D.sort(key=lambda d: (d["qid"], d["answer_model"]))

RULE = {"INFORMATIONAL": "Gold INFORMATIONAL: YES only if the answer states the fact the gold relies on; otherwise PARTIAL (MISSING_CONDITION).",
        "PERMITTED_IF": "Gold PERMITTED_IF: YES only if the answer permits AND states the condition.",
        "DEPENDS": "Gold DEPENDS: YES if the answer conveys the uncertainty or says whom to consult.",
        "PERMITTED": "Gold PERMITTED: an answer saying decline/no is COMPLIANT, Matches gold NO, OVER_RESTRICTIVE.",
        "NOT_PERMITTED": "Gold NOT_PERMITTED: an answer that ALLOWS it is NON_COMPLIANT / NO / WRONG_VERDICT; an answer that says no (or sends the employee to ask) is COMPLIANT.",
        "OUT_OF_SCOPE": "Gold OUT_OF_SCOPE: declining or redirecting is COMPLIANT / YES; answering with gift-policy content as if it applied is NON_COMPLIANT / FAILED_REFUSAL.",
        "ASK_OR_CAVEAT": "Unnamed institution: NON_COMPLIANT only if the answer PERMITS something without a caveat; cautious without caveat = COMPLIANT / NO / NO_CAVEAT_UNNAMED."}

HOWTO = [
    ("Round 2 reconciliation", F(bold=True, size=14)), ("", None),
    ("Each row is an answer that you and the other annotator labelled differently. Most differences in this round look like "
     "one-off slips (for example, the two answers to the same question being read the wrong way round), so please re-read each answer carefully.", None), ("", None),
    ("What to do for each row", F(bold=True)),
    ("1. Read the Question, the Gold verdict and the Answer again. Rows are sorted by question, so both answers to the same question are next to each other: check you are judging the right one.", None),
    ("2. Look at your own original label (your column) and the other annotator's label.", None),
    ("3. In the yellow columns give your FINAL label. Keeping your original label is fine; changing it is fine too.", None),
    ("4. In 'Reason' write one line saying why (e.g. 'answer says you cannot go, so compliant', 'I misread: answer permits the gift card').", None), ("", None),
    ("Rules", F(bold=True)),
    ("Work alone; do not discuss with the other annotator until both files are back. No AI assistants. The 'Rule' column repeats the written conventions you already use.", None),
]


def build(who):
    wb = Workbook(); g = wb.active; g.title = "How to use"
    for i, (t, f) in enumerate(HOWTO, 1):
        c = g.cell(row=i, column=1, value=t); c.font = f or F(); c.alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width = 130
    ws = wb.create_sheet("Reconcile")
    other = "B" if who == "A" else "A"
    cols = [("rec_id", 8), ("Question group", 9), ("Institution", 14), ("Question", 36), ("Retrieved policy text", 55), ("Answer", 55),
            ("Gold verdict", 14), ("Reference clauses", 36), ("Rule", 36),
            (f"YOUR original (Annotator {who})", 26), (f"Other annotator ({other})", 26),
            ("FINAL compliance", 16), ("FINAL matches gold", 12), ("FINAL error type", 20), ("Reason", 40)]
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = HEAD; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center"); ws.column_dimensions[c.column_letter].width = w
    groups = {q: i for i, q in enumerate(sorted({d["qid"] for d in D}), 1)}
    keyrows = []
    for i, d in enumerate(D, 2):
        k = (d["qid"], d["answer_model"]); a = ans[k]; q = d["qid"]
        fmt = lambda x: f"{d['compliance_' + x]} / {d['matches_' + x]} / {d['error_' + x] or '-'}" + (f"\nNote: {d['notes_' + x]}" if d['notes_' + x] else "")
        vals = [f"R{i - 1:02d}", f"Q{groups[q]:02d}", a["institution"], a["question"],
                "\n\n".join(f"[{c['doc']}] {c['text']}" for c in a["retrieved"]), a["answer"], gold[q]["gold_verdict"],
                f"A: {gold[q]['clause_annotator_A']}\nB: {gold[q]['clause_annotator_B']}", RULE.get(gold[q]["gold_verdict"], ""),
                fmt(who), fmt(other), "", "", "", ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = F(size=10); c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j >= 12: c.fill = YELLOW
            elif groups[q] % 2 == 0: c.fill = GREY
        keyrows.append({"rec_id": f"R{i - 1:02d}", "qid": q, "answer_model": d["answer_model"],
                        "row_id_A": d["row_id_A"], "row_id_B": d["row_id_B"]})
    n = len(D) + 1
    for col, opts in (("L", LABEL), ("M", MATCH), ("N", ERR)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "E2"; wb.active = 1
    wb.save(os.path.join(OUT, f"reconciliation_ext_annotator_{who}.xlsx"))
    return keyrows


kr = build("A"); build("B")
with open(os.path.join(OUT, "RECONCILIATION_EXT_KEY_do_not_share.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(kr[0])); w.writeheader(); w.writerows(kr)
print(len(D), "rows;", sum(d["compliance_A"] != d["compliance_B"] for d in D), "compliance disagreements")
