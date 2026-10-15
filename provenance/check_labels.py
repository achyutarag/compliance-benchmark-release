"""Validate completed labelling workbooks and compute agreement on the shared rows."""
import csv, os, shutil
from collections import Counter
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # repo root; L = os.path.join(HERE, "labelling"); C = os.path.join(L, "completed")
os.makedirs(C, exist_ok=True)
SRC = {"A": r"<downloads>\labelling_annotator_A_completed.xlsx",
       "B": r"<downloads>\Labelling-Annotator-B-Completed.xlsx"}
VALID = {"Compliance label": {"COMPLIANT", "NON_COMPLIANT"}, "Matches gold": {"YES", "PARTIAL", "NO"},
         "Gold clause in retrieved text": {"YES", "PARTIAL", "NO", "N/A"},
         "Error type": {"", "WRONG_VERDICT", "MISSING_CONDITION", "WRONG_INSTITUTION_RULE", "INVENTED_RULE",
                        "OVER_RESTRICTIVE", "NO_CAVEAT_UNNAMED", "FAILED_REFUSAL", "OTHER"}}
key = {r["row_id"]: r for r in csv.DictReader(open(os.path.join(L, "LABELLING_KEY_do_not_share.csv"), encoding="utf-8"))}

def read(path):
    wb = load_workbook(path, read_only=True)
    # find the labels sheet by its header (one annotator's file was re-saved as a single 'Sheet1')
    ws = next(s for s in wb.worksheets if next(s.iter_rows(values_only=True, max_row=1))[0] == "row_id")
    rows = list(ws.iter_rows(values_only=True)); hdr = [h for h in rows[0]]
    return hdr, [{h: ("" if v is None else str(v).strip()) for h, v in zip(hdr, r)} for r in rows[1:] if r[0]]

labels = {}
for who, src in SRC.items():
    dst = os.path.join(C, f"labelling_annotator_{who}_completed.xlsx"); shutil.copy2(src, dst)
    hdr0, orig = read(os.path.join(os.environ.get("ORIG_DIR", L), f"labelling_annotator_{who}.xlsx"))
    hdr, rows = read(dst)
    print(f"\n=== Annotator {who}: {len(rows)} rows (original {len(orig)}) ===")
    print("header unchanged:", hdr == hdr0)
    o = {r["row_id"]: r for r in orig}
    altered = [r["row_id"] for r in rows if any(r[c] != o[r["row_id"]][c] for c in ("Question", "Answer", "Gold verdict (reference)"))]
    print("rows missing:", sorted(set(o) - {r["row_id"] for r in rows}), "| read-only columns altered:", altered[:5])
    problems = []
    for r in rows:
        for col, ok in VALID.items():
            if col != "Error type" and not r[col]: problems.append((r["row_id"], col, "blank"))
            elif r[col] not in ok: problems.append((r["row_id"], col, r[col]))
        if r["Matches gold"] in ("PARTIAL", "NO") and not r["Error type"]:
            problems.append((r["row_id"], "Error type", "missing although Matches gold is " + r["Matches gold"]))
    print("validation problems:", len(problems), problems[:8])
    for col in ("Compliance label", "Matches gold", "Error type", "Gold clause in retrieved text"):
        print(f"  {col}: {dict(Counter(r[col] or '(blank)' for r in rows))}")
    labels[who] = {(key[r["row_id"]]["qid"], key[r["row_id"]]["answer_model"]): r for r in rows}

def kappa(pairs):
    n = len(pairs); po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n**2
    return po, (po - pe) / (1 - pe) if pe < 1 else float("nan")

shared = sorted(set(labels["A"]) & set(labels["B"]))
print(f"\n=== Agreement on {len(shared)} shared rows ===")
for col in ("Compliance label", "Matches gold", "Gold clause in retrieved text", "Error type"):
    pairs = [(labels["A"][k][col] or "-", labels["B"][k][col] or "-") for k in shared]
    po, k = kappa(pairs); print(f"  {col:30} raw {po:.1%}  kappa {k:.2f}")
dis = [k for k in shared if labels["A"][k]["Compliance label"] != labels["B"][k]["Compliance label"]
       or labels["A"][k]["Matches gold"] != labels["B"][k]["Matches gold"]]
print(f"\nrows disagreeing on compliance label or matches-gold: {len(dis)}")
with open(os.path.join(C, "shared_row_disagreements.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["qid", "answer_model", "row_id_A", "row_id_B", "compliance_A", "compliance_B", "matches_A", "matches_B",
                "error_A", "error_B", "clause_A", "clause_B", "notes_A", "notes_B"])
    for k in dis:
        a, b = labels["A"][k], labels["B"][k]
        w.writerow([k[0], k[1], a["row_id"], b["row_id"], a["Compliance label"], b["Compliance label"], a["Matches gold"],
                    b["Matches gold"], a["Error type"], b["Error type"], a["Gold clause in retrieved text"],
                    b["Gold clause in retrieved text"], a["Notes"], b["Notes"]])
        print(f"  {k[0]:4} comp {a['Compliance label'][:6]:6}/{b['Compliance label'][:6]:6}  match {a['Matches gold']:7}/{b['Matches gold']:7}"
              f"  err {a['Error type'][:14]:14}/{b['Error type'][:14]}")
