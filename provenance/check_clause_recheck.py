"""Validate the completed clause re-check: completeness, quote verification, consistency with agreed examples."""
import csv, json, os, re, shutil, sys
from collections import Counter
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # repo root; R = os.path.join(HERE, "labelling", "reconciliation")
src = sys.argv[1]
dst = os.path.join(R, "clause_recheck_completed.xlsx"); shutil.copy2(src, dst)
key = {r["row_id"]: r for r in csv.DictReader(open(os.path.join(R, "CLAUSE_RECHECK_KEY_do_not_share.csv"), encoding="utf-8"))}
ans = {}
for m in ("claude-haiku-4-5", "gpt-4o-mini"):
    for r in map(json.loads, open(os.path.join(HERE, "answers", f"answers_{m}.jsonl"), encoding="utf-8")):
        ans[(r["qid"], m)] = r

wb = load_workbook(dst, read_only=True)
ws = next(s for s in wb.worksheets if next(s.iter_rows(values_only=True, max_row=1))[0] == "row_id")
rows = list(ws.iter_rows(values_only=True)); hdr = rows[0]
ci = {h: i for i, h in enumerate(hdr)}
LAB = hdr[6]; QUO = hdr[7]; NOTE = hdr[8]
print("header:", hdr[6:9])
norm = lambda s: re.sub(r"[^a-z0-9$%]+", " ", (s or "").lower()).strip()

out, problems = [], []
for r in rows[1:]:
    if not r[0]: continue
    rid = r[0]; k = key[rid]; it = ans[(k["qid"], k["answer_model"])]
    lab = (str(r[6] or "")).strip().upper(); quote = str(r[7] or "").strip(); note = str(r[8] or "").strip()
    ctx = norm(" ".join(c["text"] for c in it["retrieved"]))
    status = "ok"
    if lab not in ("YES", "PARTIAL", "NO", "N/A"):
        status = f"invalid/blank label '{lab}'"
    elif lab in ("YES", "PARTIAL"):
        if not quote:
            status = "YES/PARTIAL without quote"
        else:
            # allow ellipses: every segment of >= 3 words must appear verbatim (after normalisation)
            segs = [norm(s) for s in re.split(r"\.\.\.|…|\[\.\.\.\]", quote)]
            segs = [s for s in segs if len(s.split()) >= 3] or [norm(quote)]
            missing = [s for s in segs if s not in ctx]
            if missing:
                status = "quote not found in retrieved text: " + missing[0][:80]
    elif quote:
        status = f"{lab} but a quote was given"
    if k["qid"][0] in "NO" and lab != "N/A":
        status = (status + "; " if status != "ok" else "") + f"{k['qid']} is unnamed/out-of-scope but labelled {lab}"
    if status != "ok": problems.append((rid, k["qid"], lab, status, quote[:90]))
    out.append({"row_id": rid, "qid": k["qid"], "answer_model": k["answer_model"], "clause_label": lab,
                "quote": quote, "notes": note, "quote_check": status})

print(f"rows: {len(out)} | labels: {dict(Counter(o['clause_label'] for o in out))}")
print(f"problems: {len(problems)}")
for p in problems: print("  ", p)
agreed = {"D09": "PARTIAL", "D10": "YES", "M02": "NO", "M06": "NO", "M08": "PARTIAL", "N04": "N/A", "O03": "N/A",
          "P07": "NO", "T13": "NO", "V01": "NO", "V02": "NO", "V10": "YES", "V11": "NO"}
# agreed examples were specific (qid, model) rows; compare on the same answers
recon = load_workbook(os.path.join(R, "label_reconciliation.xlsx"), read_only=True)["Clause examples"]
lk = {r["row_id"]: r for r in csv.DictReader(open(os.path.join(HERE, "labelling", "LABELLING_KEY_do_not_share.csv"), encoding="utf-8"))}
ex = {}
for r in list(recon.iter_rows(values_only=True))[1:]:
    a_id = r[0].split(" / ")[0]; ex[(lk[a_id]["qid"], lk[a_id]["answer_model"])] = r[9]
byk = {(o["qid"], o["answer_model"]): o for o in out}
match = [(k, v, byk[k]["clause_label"]) for k, v in ex.items()]
print(f"\nagreed reconciliation examples reproduced: {sum(a == b for _, a, b in match)}/{len(match)}")
for k, a, b in match:
    if a != b: print(f"   {k[0]} {'(model hidden)'}: agreed {a} vs re-check {b}")
with open(os.path.join(R, "clause_recheck_validated.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(out[0])); w.writeheader(); w.writerows(out)
